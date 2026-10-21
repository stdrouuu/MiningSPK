import openpyxl
wb = openpyxl.load_workbook('docs/spreadsheet/SPK_SMARTER_REAL_MHS-R008_v1-BARU.xlsx', data_only=True)
print(wb.sheetnames)
out = []
for sn in ['00_COVER','01_KRITERIA','02_PROFIL','03_ALTERNATIF','04_SMARTER','05_ARSIP_221','06_DASHBOARD']:
    ws = wb[sn]
    out.append(f"== {sn} {ws.dimensions} ==")
    for row in ws.iter_rows(values_only=True):
        vals = [(str(v)[:60] if v is not None else "") for v in row]
        if any(vals):
            out.append(" | ".join(vals))
    out.append("")
Path = __import__("pathlib").Path("docs/assets_laporan/dump_r008.txt")
Path.write_text("\n".join(out), encoding="utf-8")
print("dumped", Path, len(out))
